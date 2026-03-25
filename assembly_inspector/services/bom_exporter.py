"""BOM export and persistence helpers."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable

from assembly_inspector.logging_utils import get_logger
from assembly_inspector.services.bom_manager import BOMRow

LOGGER = get_logger(__name__)


class BOMExporter:
    """Export BOM rows to files and FreeCAD spreadsheets."""

    DEFAULT_SHEET_NAME = "AssemblyInspectorBOM"

    def export_csv(self, rows: Iterable[BOMRow], path: str | Path) -> Path:
        """Export BOM rows to CSV."""
        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)

        serialized = [row.to_dict() for row in rows]
        headers = self._resolve_headers(serialized)

        with target.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers)
            writer.writeheader()
            writer.writerows(serialized)

        LOGGER.info("BOM exported to CSV: %s", target)
        return target

    def export_xlsx(self, rows: Iterable[BOMRow], path: str | Path) -> Path:
        """Export BOM rows to XLSX.

        TODO: If the project standardizes an XLSX dependency strategy, this can
        move from optional import to required runtime dependency.
        """
        try:
            from openpyxl import Workbook  # type: ignore
        except ImportError as exc:
            raise RuntimeError("XLSX export requires 'openpyxl' to be installed") from exc

        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)

        serialized = [row.to_dict() for row in rows]
        headers = self._resolve_headers(serialized)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "BOM"
        sheet.append(headers)
        for item in serialized:
            sheet.append([item.get(header, "") for header in headers])
        workbook.save(target)

        LOGGER.info("BOM exported to XLSX: %s", target)
        return target

    def write_to_document_spreadsheet(
        self,
        document: Any,
        rows: Iterable[BOMRow],
        sheet_name: str = DEFAULT_SHEET_NAME,
    ) -> bool:
        """Persist BOM rows into a Spreadsheet::Sheet object.

        Returns ``True`` when write appears successful.
        """
        if document is None:
            return False

        sheet = self._get_or_create_sheet(document, sheet_name)
        if sheet is None or not hasattr(sheet, "set"):
            return False

        serialized = [row.to_dict() for row in rows]
        headers = self._resolve_headers(serialized)

        self._write_sheet_rows(sheet, headers, serialized)

        # Recompute may not exist on all test doubles or in all contexts.
        try:
            if hasattr(document, "recompute"):
                document.recompute()
        except Exception:
            LOGGER.exception("Document recompute failed after BOM sheet write")
        return True

    def _get_or_create_sheet(self, document: Any, sheet_name: str) -> Any | None:
        try:
            sheet = document.getObject(sheet_name) if hasattr(document, "getObject") else None
            if sheet is not None:
                return sheet
            if hasattr(document, "addObject"):
                return document.addObject("Spreadsheet::Sheet", sheet_name)
            return None
        except Exception:
            LOGGER.exception("Failed to create/access BOM spreadsheet '%s'", sheet_name)
            return None

    def _write_sheet_rows(
        self,
        sheet: Any,
        headers: list[str],
        rows: list[dict[str, str]],
    ) -> None:
        for col_idx, header in enumerate(headers, start=1):
            sheet.set(f"{self._to_column(col_idx)}1", header)

        row_idx = 2
        for item in rows:
            for col_idx, header in enumerate(headers, start=1):
                sheet.set(f"{self._to_column(col_idx)}{row_idx}", str(item.get(header, "")))
            row_idx += 1

        # Clear trailing area from previous writes up to a conservative window.
        # TODO: replace with exact used-range clearing when robust API support is confirmed.
        for clear_row in range(row_idx, row_idx + 100):
            for col_idx in range(1, len(headers) + 1):
                sheet.set(f"{self._to_column(col_idx)}{clear_row}", "")

    @staticmethod
    def _resolve_headers(rows: list[dict[str, str]]) -> list[str]:
        base = ["Item", "PartId", "Label", "Quantity", "Document"]
        extra: set[str] = set()
        for row in rows:
            for key in row.keys():
                if key not in base:
                    extra.add(key)
        return base + sorted(extra)

    @staticmethod
    def _to_column(index: int) -> str:
        if index < 1:
            raise ValueError("index must be >= 1")
        chars: list[str] = []
        value = index
        while value > 0:
            value, remainder = divmod(value - 1, 26)
            chars.append(chr(65 + remainder))
        return "".join(reversed(chars))

